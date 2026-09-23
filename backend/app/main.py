from __future__ import annotations

import difflib
import json
import logging
import os
import re
import sqlite3
import zipfile
from contextlib import contextmanager
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field

ROOT = Path(__file__).resolve().parents[2]
DATA_FILE = ROOT / "data" / "catalog.json"


def load_env() -> None:
    for env_path in (ROOT / ".env", ROOT / "backend" / ".env"):
        if not env_path.exists():
            continue
        for line in env_path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            os.environ.setdefault(key.strip(), value.strip().strip('"\''))


load_env()
logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s %(message)s")
log = logging.getLogger("ekt-assistant")
DEMO_MODE = os.getenv("DEMO_MODE", "true").lower() in {"1", "true", "yes"}
DB_FILE = Path(os.getenv("DATABASE_PATH", str(ROOT / "backend" / "data" / "ekt_assistant.sqlite3")))
if not DB_FILE.is_absolute():
    DB_FILE = ROOT / DB_FILE


def load_catalog() -> list[dict[str, Any]]:
    if not DATA_FILE.exists():
        raise RuntimeError("Demo catalog missing. Run: node scripts/prepare_catalog.mjs")
    return json.loads(DATA_FILE.read_text(encoding="utf-8"))["products"]


PRODUCTS = load_catalog()
BY_ID = {str(p["id"]): p for p in PRODUCTS}
SYSTEM_PROMPT = (Path(__file__).with_name("system_prompt.txt").read_text(encoding="utf-8").strip()
                 if Path(__file__).with_name("system_prompt.txt").exists() else "Ты консультант ekt.kz. Используй только данные backend tools.")


@contextmanager
def db() -> Any:
    DB_FILE.parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(DB_FILE, timeout=10)
    con.row_factory = sqlite3.Row
    try:
        yield con
        con.commit()
    finally:
        con.close()


def init_db() -> None:
    with db() as con:
        con.executescript("""
            PRAGMA journal_mode=WAL;
            CREATE TABLE IF NOT EXISTS sessions (
                id TEXT PRIMARY KEY, state_json TEXT NOT NULL DEFAULT '{}', updated_at TEXT DEFAULT CURRENT_TIMESTAMP
            );
            CREATE TABLE IF NOT EXISTS cart_items (
                product_id TEXT PRIMARY KEY, quantity INTEGER NOT NULL CHECK(quantity > 0), updated_at TEXT DEFAULT CURRENT_TIMESTAMP
            );
            CREATE TABLE IF NOT EXISTS cart_actions (
                id INTEGER PRIMARY KEY AUTOINCREMENT, product_id TEXT NOT NULL, quantity INTEGER NOT NULL,
                confirmed INTEGER NOT NULL, created_at TEXT DEFAULT CURRENT_TIMESTAMP
            );
        """)


def norm(value: Any) -> str:
    value = str(value or "").lower().replace("ё", "е")
    return re.sub(r"[^a-zа-я0-9]+", "", value)


BY_ARTICLE = {norm(p.get("article")): p for p in PRODUCTS if p.get("article")}


def tokens(value: str) -> list[str]:
    return [t for t in re.findall(r"[a-zа-я0-9]+", value.lower().replace("ё", "е")) if len(t) > 1]


def search_products(query: str, limit: int = 8) -> list[dict[str, Any]]:
    raw = query.strip()
    key = norm(raw)
    if not key:
        return []
    exact = BY_ARTICLE.get(key)
    if exact:
        return [exact]
    words = tokens(raw)
    # User language often calls Legrand simply "лг"; treat this as a brand hint.
    alias_brand = any(w in {"лг", "лгранд", "legrand"} for w in words)
    scored: list[tuple[float, dict[str, Any]]] = []
    for p in PRODUCTS:
        name = norm(p.get("name"))
        brand = norm(p.get("brand"))
        article = norm(p.get("article"))
        erp_article = norm(p.get("erp_article"))
        if key and (key == erp_article or key in name):
            score = 10.0
        else:
            score = 0.0
            if alias_brand and brand == "legrand":
                score += 1.8
            for word in words:
                wn = norm(word)
                if not wn:
                    continue
                if wn == article or wn == erp_article:
                    score += 12
                elif wn in name:
                    score += 3.2 if wn in brand else 2.2
                else:
                    similarity = max(difflib.SequenceMatcher(None, wn, token).ratio() for token in tokens(p.get("name", "")) or [""])
                    if similarity >= 0.78:
                        score += 0.8
            if words and all(w in {"найди", "найти", "покажи", "товар", "есть", "ли", "цена", "характеристики", "какая", "какие", "автомат"} for w in words):
                if alias_brand and brand == "legrand":
                    score += 1.0
        if score > 0:
            scored.append((score, p))
    scored.sort(key=lambda row: (-row[0], int(row[1]["id"])))
    return [p for _, p in scored[:limit]]


def get_product(product_id: str) -> dict[str, Any] | None:
    return BY_ID.get(str(product_id)) or BY_ARTICLE.get(norm(product_id))


def city_stock(product: dict[str, Any], city: str | None = None) -> dict[str, Any]:
    if not city:
        return {"known": product.get("stock") is not None, "quantity": product.get("stock"), "city": None}
    stores = product.get("stores")
    if stores is None:
        return {"known": False, "quantity": None, "city": city}
    aliases = {"алмааты": "алматы", "алмаата": "алматы", "астана": "нурсултан", "нурсултан": "нурсултан"}
    city_key = aliases.get(norm(city), norm(city))
    matches = [s for s in stores if city_key and city_key in norm(s["name"])]
    if not matches:
        return {"known": False, "quantity": None, "city": city}
    return {"known": True, "quantity": sum(int(s["quantity"]) for s in matches), "city": city, "stores": matches}


def analogs(product: dict[str, Any], limit: int = 4) -> list[dict[str, Any]]:
    props = product.get("properties") or {}
    target_series = (props.get("series") or "").upper().replace(" ", "")
    if not target_series:
        return []
    result: list[tuple[int, dict[str, Any]]] = []
    for other in PRODUCTS:
        if other["id"] == product["id"]:
            continue
        op = other.get("properties") or {}
        if (op.get("series") or "").upper().replace(" ", "") != target_series:
            continue
        same = sum(1 for field in ("poles", "amperage", "breaking_capacity") if props.get(field) and props.get(field) == op.get(field))
        # Same product family and poles are minimum requirements for this demo analog list.
        if same >= 1 and props.get("poles") == op.get("poles"):
            result.append((same, other))
    result.sort(key=lambda row: (-row[0], int(row[1]["id"])))
    return [p for _, p in result[:limit]]


def compare_analog(base: dict[str, Any], candidate: dict[str, Any]) -> dict[str, Any]:
    names = {"poles": "число полюсов", "amperage": "номинальный ток", "breaking_capacity": "отключающая способность"}
    a, b = base.get("properties") or {}, candidate.get("properties") or {}
    matching, differences = [], []
    for field, label in names.items():
        if a.get(field) and b.get(field):
            (matching if a[field] == b[field] else differences).append(f"{label}: {a[field]}{'А' if field == 'amperage' else ' кА' if field == 'breaking_capacity' else ''} / {b[field]}{'А' if field == 'amperage' else ' кА' if field == 'breaking_capacity' else ''}")
    return {"matching": matching, "differences": differences}


def get_purchase_conditions() -> dict[str, Any]:
    # The supplied materials do not contain payment, delivery, or minimum-order terms.
    return {"available": False, "payment": None, "delivery": None, "minimum_order": None,
            "message": "В загруженных материалах не указаны условия оплаты, доставки и минимальной партии. Для точного ответа обратитесь к менеджеру ekt.kz."}


class ChatRequest(BaseModel):
    session_id: str = Field(min_length=4, max_length=100)
    message: str = Field(min_length=1, max_length=2000)


class CartRequest(BaseModel):
    product_id: str
    quantity: int = Field(gt=0, le=10000)
    confirmed: bool = False


def load_state(session_id: str) -> dict[str, Any]:
    with db() as con:
        row = con.execute("SELECT state_json FROM sessions WHERE id=?", (session_id,)).fetchone()
    return json.loads(row[0]) if row else {}


def save_state(session_id: str, state: dict[str, Any]) -> None:
    with db() as con:
        con.execute("INSERT INTO sessions(id,state_json,updated_at) VALUES(?,?,CURRENT_TIMESTAMP) ON CONFLICT(id) DO UPDATE SET state_json=excluded.state_json,updated_at=CURRENT_TIMESTAMP", (session_id, json.dumps(state, ensure_ascii=False)))


def cart_contents() -> dict[str, Any]:
    with db() as con:
        rows = con.execute("SELECT product_id,quantity FROM cart_items ORDER BY updated_at DESC").fetchall()
    items, total = [], 0
    for row in rows:
        p = get_product(row["product_id"])
        if not p:
            continue
        line_total = int(p["price"]) * int(row["quantity"])
        total += line_total
        items.append({"product": p, "quantity": row["quantity"], "line_total": line_total})
    return {"items": items, "total": total}


def add_to_cart(product_id: str, quantity: int, confirmed: bool) -> dict[str, Any]:
    # Repeat catalog/availability validation immediately before the cart write.
    product = get_product(product_id)
    if not product:
        raise HTTPException(404, "Товар не найден в каталоге.")
    if not confirmed:
        raise HTTPException(409, "Сначала подтвердите добавление товара.")
    if quantity <= 0:
        raise HTTPException(422, "Количество должно быть больше нуля.")
    current = product.get("stock")
    if current is None:
        raise HTTPException(409, "Актуальный остаток этого товара неизвестен. Добавление отменено.")
    with db() as con:
        con.execute("BEGIN IMMEDIATE")
        existing = con.execute("SELECT quantity FROM cart_items WHERE product_id=?", (product["id"],)).fetchone()
        already_in_cart = int(existing[0]) if existing else 0
        available = max(0, int(current) - already_in_cart)
        if quantity > available:
            raise HTTPException(409, f"Доступно только {available} шт.")
        con.execute("INSERT INTO cart_items(product_id,quantity) VALUES(?,?) ON CONFLICT(product_id) DO UPDATE SET quantity=quantity+excluded.quantity,updated_at=CURRENT_TIMESTAMP", (product["id"], quantity))
        con.execute("INSERT INTO cart_actions(product_id,quantity,confirmed) VALUES(?,?,1)", (product["id"], quantity))
    return {"ok": True, "product": product, "quantity_added": quantity, "cart": cart_contents()}


def extract_quantity(message: str) -> int | None:
    match = re.search(r"\b(\d{1,5})\s*(?:шт\.?|штук[аи]?|единиц[уы]?)?\b", message.lower())
    return int(match.group(1)) if match else None


def yes_intent(message: str) -> bool:
    return bool(re.search(r"\b(да|добавь|добавить|подтверждаю|подтверждаю|согласен|ок|конечно)\b", message.lower()))


def no_intent(message: str) -> bool:
    return bool(re.search(r"\b(нет|отмена|отмени|не надо|не добавляй)\b", message.lower()))


def product_card(p: dict[str, Any]) -> dict[str, Any]:
    return {k: p.get(k) for k in ("id", "name", "article", "price", "image", "url", "brand", "stock", "properties")}


def price_text(price: int | None) -> str:
    return f"{int(price):,}".replace(",", " ") + " ₸" if price is not None else "цена не указана"


def local_chat(message: str, state: dict[str, Any]) -> tuple[str, list[dict[str, Any]], dict[str, Any]]:
    text = message.strip()
    lowered = text.lower().replace("ё", "е")
    cards: list[dict[str, Any]] = []
    selected = get_product(str(state.get("selected_id", ""))) if state.get("selected_id") else None
    pending = state.get("pending")

    if pending and yes_intent(text):
        p = get_product(str(pending["product_id"]))
        quantity = int(pending["quantity"])
        try:
            result = add_to_cart(str(pending["product_id"]), quantity, confirmed=True)
            state.pop("pending", None)
            state["selected_id"] = p["id"]
            return f"Готово, добавил {quantity} шт. «{p['name']}» в корзину. Можно перейти в корзину.", [product_card(p)], state | {"cart": result["cart"]}
        except HTTPException as exc:
            state.pop("pending", None)
            return str(exc.detail), [product_card(p)] if p else [], state

    if pending and no_intent(text):
        state.pop("pending", None)
        return "Хорошо, товар не добавлял. Если понадобится — скажите количество.", [], state

    qty = extract_quantity(text)
    asks_quantity = qty is not None and bool(re.search(r"(нужн|надо|хочу|добав|штук|шт\b|количеств)", lowered))
    if asks_quantity:
        requested_product = selected
        article_match = re.search(r"\b\d{4,}\b", text)
        if article_match:
            requested_product = get_product(article_match.group(0)) or requested_product
        if not requested_product:
            return "Сначала выберите товар из каталога, тогда проверю остаток и подготовлю добавление.", [], state
        state["selected_id"] = requested_product["id"]
        stock = requested_product.get("stock")
        if stock is None:
            return f"Для «{requested_product['name']}» в демо-каталоге остаток не указан, поэтому я не могу подтвердить количество для корзины.", [product_card(requested_product)], state
        with db() as con:
            in_cart_row = con.execute("SELECT quantity FROM cart_items WHERE product_id=?", (requested_product["id"],)).fetchone()
        available = max(0, int(stock) - (int(in_cart_row[0]) if in_cart_row else 0))
        if available < 1:
            return "Сейчас доступного остатка нет, добавить товар нельзя.", [product_card(requested_product)], state
        if qty > available:
            state["pending"] = {"product_id": requested_product["id"], "quantity": available}
            return f"В наличии только {available} шт. Добавить {available} в корзину?", [product_card(requested_product)], state
        if qty < 1:
            return "Количество должно быть больше нуля.", [product_card(requested_product)], state
        state["pending"] = {"product_id": requested_product["id"], "quantity": qty}
        return f"Общий остаток {stock} шт., ещё можно добавить {available}. Добавить {qty} шт. «{requested_product['name']}» в корзину?", [product_card(requested_product)], state

    is_analog = any(k in lowered for k in ("аналог", "замен", "похож"))
    if is_analog:
        if not selected:
            found = search_products(text)
            selected = found[0] if found else None
        if not selected:
            return "Сначала найдите товар, для которого нужно подобрать аналог.", [], state
        found = analogs(selected)
        cards = [product_card(p) for p in found]
        state["selected_id"] = selected["id"]
        state["last_product_ids"] = [p["id"] for p in found]
        if not found:
            return "В загруженном каталоге не нашёл достаточно близких аналогов с подтверждёнными характеристиками.", [], state
        lines = []
        for p in found:
            comparison = compare_analog(selected, p)
            same = ", ".join(comparison["matching"]) or "та же серия и категория"
            diff = ", ".join(comparison["differences"]) or "различия в известных параметрах не указаны"
            stock_text = f"{p['stock']} шт." if p.get("stock") is not None else "остаток в демо-данных не указан"
            lines.append(f"• {p['name']} · {p['article'] or 'артикул не указан'} · {price_text(p['price'])} · {stock_text}. Совпадает: {same}. Отличается: {diff}.")
        return "Похожие варианты из той же серии (сверьте требования по току):\n" + "\n".join(lines), cards, state

    if any(k in lowered for k in ("оплат", "доставк", "минимальн", "условия покупки", "как купить")):
        return get_purchase_conditions()["message"], [], state

    characteristic = any(k in lowered for k in ("характерист", "отключающ", "полюс", "напряжен", "ток"))
    certificate_question = "сертификат" in lowered
    relative = re.search(r"\b(перв(ый|ого)|втор(ой|ого)|треть(его|ему)|последн(ий|его))\b", lowered)
    relative_product = None
    if relative and state.get("last_product_ids"):
        position = 0 if relative.group(1).startswith("перв") else 1 if relative.group(1).startswith("втор") else 2 if relative.group(1).startswith("трет") else -1
        ids = state["last_product_ids"]
        relative_product = get_product(ids[position]) if len(ids) > abs(position) - (1 if position < 0 else 0) else None
    query = text
    query = re.sub(r"^(найди|найти|покажи|есть ли|какая цена у|цена у|какие|какая)\s*", "", query, flags=re.I)
    query = re.sub(r"\b(в наличии|по городу|характеристики|характеристик|цена|есть ли)\b", "", query, flags=re.I).strip(" ?!. ,")
    direct = search_products(query)
    if relative_product:
        direct = [relative_product]
    article = re.search(r"\b\d{4,}\b", text)
    if article:
        direct = search_products(article.group(0)) or direct
    if not direct and selected:
        direct = [selected]
    city_match = re.search(r"\b(?:в|по)\s+(алматы|алмааты|астана|нур-султан|шымкент|тараз|атырау|караганда|актау)\b", lowered)
    price_question = "цен" in lowered or "сколько стоит" in lowered
    stock_question = bool(re.search(r"(есть ли|налич|в\s+алматы|в\s+шымкенте|остаток)", lowered))
    if not direct:
        return "Не нашёл подходящего товара в загруженных данных. Попробуйте указать артикул или часть названия.", [], state

    state["last_product_ids"] = [p["id"] for p in direct]
    state["selected_id"] = direct[0]["id"]
    cards = [product_card(p) for p in direct[:6]]
    if characteristic:
        p = direct[0]
        props = p.get("properties") or {}
        values = []
        for field, label, suffix in (("poles", "Полюсов", ""), ("amperage", "Номинальный ток", " А"), ("breaking_capacity", "Отключающая способность", " кА"), ("voltage", "Напряжение", ""), ("installation", "Монтаж", "")):
            if props.get(field): values.append(f"{label}: {props[field]}{suffix}")
        if not values and p.get("description"):
            return f"Характеристики «{p['name']}»:\n{p['description']}", cards[:1], state
        return f"Характеристики «{p['name']}»:\n" + ("\n".join(f"• {v}" for v in values) if values else "В демо-данных характеристики не указаны."), cards[:1], state
    if certificate_question:
        p = direct[0]
        if p.get("certificate_url"):
            return f"Сертификат «{p['name']}»: {p['certificate_url']}", cards[:1], state
        return f"В загруженных данных для «{p['name']}» ссылка или файл сертификата не указаны.", cards[:1], state
    if city_match or stock_question:
        p = direct[0]
        stock = city_stock(p, city_match.group(1) if city_match else None)
        if stock["known"]:
            where = f" в {stock['city']}" if stock.get("city") else " по всем складам"
            answer = f"Да, «{p['name']}» ({p['article']}){where}: {stock['quantity']} шт."
        else:
            answer = f"Для «{p['name']}» в загруженных демо-данных остаток{(' в ' + stock['city']) if stock.get('city') else ''} не указан."
        if p.get("price") is not None: answer += f" Цена: {price_text(p['price'])}."
        return answer, cards[:1], state
    if price_question:
        p = direct[0]
        return f"Цена «{p['name']}» ({p['article'] or 'артикул не указан'}): {price_text(p['price'])}.", cards[:1], state

    if len(direct) == 1:
        p = direct[0]
        available = f"Общий остаток: {p['stock']} шт." if p.get("stock") is not None else "Общий остаток в демо-данных не указан."
        return f"Нашёл товар: {p['name']}\nАртикул: {p['article'] or 'не указан'} · Цена: {price_text(p['price'])}\n{available}", cards, state
    lines = []
    for p in direct[:6]:
        quantity = f" · остаток {p['stock']} шт." if p.get("stock") is not None else " · остаток не указан"
        lines.append(f"• {p['name']} · {p['article'] or 'артикул не указан'} · {price_text(p['price'])}{quantity}")
    return f"Нашёл {len(direct)} подходящих товаров:\n" + "\n".join(lines) + "\nВыберите карточку, чтобы уточнить характеристики или остаток.", cards, state


def openai_tool_answer(message: str, state: dict[str, Any], session_id: str) -> tuple[str, list[dict[str, Any]]] | None:
    api_key = os.getenv("OPENAI_API_KEY", "").strip()
    if not api_key:
        return None
    try:
        from openai import OpenAI
        client = OpenAI(api_key=api_key, timeout=12.0, max_retries=1)
        tools = [
            {"type": "function", "function": {"name": "search_products", "description": "Ищет товары в локальном каталоге ekt.kz. Остатки неизвестны, если поле stock пустое.", "parameters": {"type": "object", "properties": {"query": {"type": "string"}}, "required": ["query"], "additionalProperties": False}}},
            {"type": "function", "function": {"name": "get_product", "description": "Возвращает достоверные цену и характеристики товара по артикулу или id.", "parameters": {"type": "object", "properties": {"product_id": {"type": "string"}}, "required": ["product_id"], "additionalProperties": False}}},
            {"type": "function", "function": {"name": "get_stock", "description": "Возвращает только сохранённый остаток товара, при наличии — для города.", "parameters": {"type": "object", "properties": {"product_id": {"type": "string"}, "city": {"type": ["string", "null"]}}, "required": ["product_id"], "additionalProperties": False}}},
            {"type": "function", "function": {"name": "find_analogs", "description": "Находит похожие позиции только в той же товарной серии и сравнивает известные характеристики.", "parameters": {"type": "object", "properties": {"product_id": {"type": "string"}}, "required": ["product_id"], "additionalProperties": False}}},
            {"type": "function", "function": {"name": "get_cart", "description": "Показывает текущую корзину.", "parameters": {"type": "object", "properties": {}, "additionalProperties": False}}},
            {"type": "function", "function": {"name": "get_purchase_conditions", "description": "Возвращает только условия оплаты, доставки и минимального заказа, если они загружены; иначе явно сообщает, что данных нет.", "parameters": {"type": "object", "properties": {}, "additionalProperties": False}}},
            {"type": "function", "function": {"name": "add_to_cart", "description": "Добавляет товар только при явном подтверждении в текущем сообщении и при наличии незакрытого запроса подтверждения в сессии.", "parameters": {"type": "object", "properties": {"product_id": {"type": "string"}, "quantity": {"type": "integer"}}, "required": ["product_id", "quantity"], "additionalProperties": False}}},
        ]
        state_context = json.dumps({"selected_id": state.get("selected_id"), "last_product_ids": state.get("last_product_ids"), "pending_confirmation": state.get("pending")}, ensure_ascii=False)
        messages: list[dict[str, Any]] = [
            {"role": "system", "content": SYSTEM_PROMPT + "\nКонтекст сессии: " + state_context},
            {"role": "user", "content": message},
        ]
        cards: list[dict[str, Any]] = []
        for _ in range(4):
            response = client.chat.completions.create(model=os.getenv("OPENAI_MODEL", "gpt-4o-mini"), messages=messages, tools=tools, tool_choice="auto")
            choice = response.choices[0].message
            if not choice.tool_calls:
                return (choice.content or "", cards) if choice.content else None
            messages.append(choice.model_dump(exclude_none=True))
            for call in choice.tool_calls:
                args = json.loads(call.function.arguments or "{}")
                name = call.function.name
                data: Any
                if name == "search_products":
                    data = search_products(args["query"])
                    cards.extend(product_card(p) for p in data[:6])
                elif name == "get_product":
                    data = get_product(args["product_id"])
                    if data: cards.append(product_card(data))
                elif name == "get_stock":
                    p = get_product(args["product_id"])
                    data = city_stock(p, args.get("city")) if p else {"known": False, "error": "Товар не найден"}
                elif name == "find_analogs":
                    p = get_product(args["product_id"])
                    data = [{"product": product_card(x), "comparison": compare_analog(p, x)} for x in analogs(p)] if p else []
                    cards.extend(row["product"] for row in data)
                elif name == "get_cart":
                    data = cart_contents()
                elif name == "get_purchase_conditions":
                    data = get_purchase_conditions()
                elif name == "add_to_cart":
                    pending = state.get("pending")
                    if yes_intent(message) and pending and str(pending.get("product_id")) == str(args.get("product_id")) and int(pending.get("quantity", 0)) == int(args.get("quantity", 0)):
                        data = add_to_cart(str(args["product_id"]), int(args["quantity"]), True)
                        state.pop("pending", None)
                    else:
                        data = {"ok": False, "message": "Корзина не изменена: нет ожидающего подтверждения пользователя."}
                else:
                    data = {"error": "Unknown tool"}
                messages.append({"role": "tool", "tool_call_id": call.id, "content": json.dumps(data, ensure_ascii=False, default=str)})
        return None
    except Exception:
        log.exception("OpenAI tool call failed; using deterministic local assistant")
        return None


app = FastAPI(title="EKT AI Shop Assistant", version="1.0.0")
app.add_middleware(CORSMiddleware, allow_origins=[os.getenv("FRONTEND_ORIGIN", "http://localhost:3000")], allow_credentials=True, allow_methods=["*"], allow_headers=["*"])


@app.on_event("startup")
def startup() -> None:
    init_db()
    log.info("Assistant ready: demo_mode=%s products=%s", DEMO_MODE, len(PRODUCTS))


@app.get("/api/health")
def health() -> dict[str, Any]:
    return {"ok": True, "demo_mode": DEMO_MODE, "products": len(PRODUCTS), "ai_enabled": bool(os.getenv("OPENAI_API_KEY"))}


@app.get("/api/products/search")
def api_search(q: str = "", limit: int = 8) -> dict[str, Any]:
    return {"items": [product_card(p) for p in search_products(q, min(max(limit, 1), 20))]}


@app.get("/api/products/{product_id}")
def api_product(product_id: str) -> dict[str, Any]:
    p = get_product(product_id)
    if not p: raise HTTPException(404, "Товар не найден")
    return p


@app.get("/api/products/{product_id}/stock")
def api_stock(product_id: str, city: str | None = None) -> dict[str, Any]:
    p = get_product(product_id)
    if not p: raise HTTPException(404, "Товар не найден")
    return city_stock(p, city)


@app.get("/api/products/{product_id}/analogs")
def api_analogs(product_id: str) -> dict[str, Any]:
    p = get_product(product_id)
    if not p: raise HTTPException(404, "Товар не найден")
    return {"items": [{"product": product_card(x), "comparison": compare_analog(p, x)} for x in analogs(p)]}


@app.get("/api/cart")
def api_cart() -> dict[str, Any]:
    return cart_contents()


@app.get("/api/purchase-conditions")
def api_purchase_conditions() -> dict[str, Any]:
    return get_purchase_conditions()


@app.post("/api/cart/add")
def api_cart_add(payload: CartRequest) -> dict[str, Any]:
    return add_to_cart(payload.product_id, payload.quantity, payload.confirmed)


@app.post("/api/chat")
def chat(payload: ChatRequest) -> dict[str, Any]:
    state = load_state(payload.session_id)
    answer: str
    cards: list[dict[str, Any]] = []
    # Cart actions use a local confirmation state machine; a model cannot bypass it.
    if state.get("pending") and (yes_intent(payload.message) or no_intent(payload.message) or extract_quantity(payload.message) is not None):
        answer, cards, state = local_chat(payload.message, state)
    else:
        model_result = None
        # Deterministic demo flows remain available offline; OpenAI tools handle open-ended requests.
        known_intent = any(word in payload.message.lower() for word in ("найди", "покажи", "аналог", "характерист", "отключающ", "цена", "сколько стоит", "есть ли", "налич", "алматы", "штук", "шт", "оплат", "доставк", "минимальн", "условия")) or bool(re.search(r"\b\d{4,}\b", payload.message))
        if not known_intent:
            model_result = openai_tool_answer(payload.message, state, payload.session_id)
        if model_result:
            answer, cards = model_result
        else:
            answer, cards, state = local_chat(payload.message, state)
    save_state(payload.session_id, state)
    return {"answer": answer, "products": cards, "session_id": payload.session_id, "cart": cart_contents(), "pending_confirmation": state.get("pending")}


def extract_docx(data: bytes) -> str:
    import io
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    return "\n".join("".join(t.text or "" for t in p.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")) for p in root.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"))


@app.post("/api/upload")
async def upload(file: UploadFile = File(...)) -> dict[str, Any]:
    filename = file.filename or "upload"
    suffix = Path(filename).suffix.lower()
    data = await file.read()
    if len(data) > 12 * 1024 * 1024: raise HTTPException(413, "Максимальный размер файла — 12 МБ")
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader
            import io
            text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
        elif suffix == ".docx": text = extract_docx(data)
        elif suffix == ".xlsx":
            from openpyxl import load_workbook
            import io
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            text = "\n".join(" ".join(str(cell) for cell in row if cell is not None) for sheet in book for row in sheet.iter_rows(values_only=True))
        elif suffix in {".png", ".jpg", ".jpeg"}:
            if not os.getenv("OPENAI_API_KEY"):
                return {"filename": filename, "message": "Файл принят. Для распознавания изображения добавьте OPENAI_API_KEY в .env; поиск в локальном каталоге работает без ключа.", "products": []}
            import base64
            from openai import OpenAI
            mime = "image/png" if suffix == ".png" else "image/jpeg"
            response = OpenAI(timeout=20, max_retries=1).chat.completions.create(model=os.getenv("OPENAI_VISION_MODEL", os.getenv("OPENAI_MODEL", "gpt-4o-mini")), messages=[{"role":"user","content":[{"type":"text","text":"Определи вероятный артикул, производителя и тип товара с изображения. Если не уверен, укажи неопределённость. Ответь кратко на русском."},{"type":"image_url","image_url":{"url":f"data:{mime};base64,{base64.b64encode(data).decode()}"}}]}], max_tokens=180)
            text = response.choices[0].message.content or ""
        else:
            raise HTTPException(415, "Поддерживаются PDF, DOCX, XLSX, PNG и JPG")
    except HTTPException: raise
    except Exception as exc:
        log.warning("Upload parse failed for %s: %s", filename, type(exc).__name__)
        raise HTTPException(422, "Не удалось прочитать файл. Проверьте его формат и попробуйте ещё раз.")
    # Uploaded source documents may contain credentials; never return those in extracted previews.
    text = re.sub(r"(?i)((?:пароль|password|api[_ -]?key|secret)\s*[:=]\s*)[^\s;,]+", r"\1[скрыто]", text)
    candidates: dict[str, dict[str, Any]] = {}
    for code in re.findall(r"\b[A-Za-zА-Яа-я0-9][A-Za-zА-Яа-я0-9-]{3,}\b", text):
        for p in search_products(code, 3): candidates[p["id"]] = p
    if not candidates:
        for p in search_products(text[:500], 4): candidates[p["id"]] = p
    return {"filename": filename, "extracted_text": text[:1800], "products": [product_card(p) for p in list(candidates.values())[:8]], "message": f"Из файла извлечён текст ({len(text)} символов). Найдено совпадений в каталоге: {len(candidates)}." if candidates else "Текст извлечён, но совпадений в локальном каталоге нет."}
